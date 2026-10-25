"""Script to generate a comprehensive, professionally styled Word Document (.docx)
Installation & Deployment Guide for the Operation Advisor system.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding of a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE", callout_type="NOTE"):
    """Add a shaded callout alert box with colored left border."""
    border_color = "0A2540" # Navy default
    bg_color = "F1F5F9"     # Light slate
    title_color = RGBColor(10, 37, 64)

    if callout_type == "IMPORTANT":
        border_color = "B91C1C" # Crimson
        bg_color = "FEF2F2"
        title_color = RGBColor(185, 28, 28)
    elif callout_type == "TIP":
        border_color = "0D9488" # Teal
        bg_color = "F0FDFA"
        title_color = RGBColor(13, 148, 136)
    elif callout_type == "WARNING":
        border_color = "D97706" # Amber
        bg_color = "FFFBEB"
        title_color = RGBColor(217, 119, 6)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = title_color

    run_text = p.add_run(text)
    run_text.font.name = "Segoe UI"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    """Add a styled code block with terminal background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "0F172A") # Dark slate terminal
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(56, 189, 248) # Cyan/Sky blue

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(tbl, col_widths, headers, data_rows, header_bg="0A2540"):
    """Format a table with industrial styling, alternating row shading, and proper padding."""
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(data_rows):
        row = tbl.add_row()
        bg_color = "F8FAFC" if (row_idx % 2 == 1) else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = Inches(col_widths[col_idx])
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(str(val))
            run.font.name = "Segoe UI"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 41, 59)

            # Set thin horizontal grid line
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                f'<w:left w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)

    doc_p = tbl._element.getparent()
    p_after = docx.text.paragraph.Paragraph(parse_xml(f'<w:p {nsdecls("w")}/>'), doc_p)
    p_after.paragraph_format.space_after = Pt(8)

def build_installation_guide():
    doc = Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)

    # ==========================================
    # COVER / TITLE PAGE HEADER
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(4)
    run_sup = title_p.add_run("CPF TRAIN 1 AUTOMATION & SAFETY SYSTEM\n")
    run_sup.font.size = Pt(11)
    run_sup.font.bold = True
    run_sup.font.color.rgb = RGBColor(14, 116, 144) # Cyan 700

    run_title = title_p.add_run("OPERATION ADVISOR")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(10, 37, 64) # Deep Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run("System Installation, Environment Setup & Deployment Guide")
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Table
    meta_tbl = doc.add_table(rows=1, cols=4)
    meta_widths = [1.6, 1.6, 1.6, 1.7]
    meta_headers = ["Document ID", "Release Version", "Target Platform", "Runtime Environment"]
    meta_data = [
        ["INS-OA-TR1-2026", "Version 4.2 Production", "Windows / Linux / macOS", "Python 3.10 - 3.13"]
    ]
    format_table(meta_tbl, meta_widths, meta_headers, meta_data, header_bg="1E293B")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_callout(
        doc,
        "Zero-Dependency Runtime Guarantee: Operation Advisor is engineered with pure-Python fallbacks. "
        "The system runs fully offline without requiring an internet connection or external cloud API keys. "
        "The PDF generator, dynamic physics simulation, forward/backward chaining, and deterministic RAG synthesizer "
        "function out of the box with zero external SaaS dependencies.",
        title="DEPLOYMENT PHILOSOPHY",
        callout_type="TIP"
    )

    # ==========================================
    # SECTION 1: HARDWARE & SYSTEM REQUIREMENTS
    # ==========================================
    h1 = doc.add_heading("1. System Pre-requisites & Minimum Specifications", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Before deploying Operation Advisor in a control room workstation or engineering server, "
        "ensure that the target host meets or exceeds the following hardware and operating system specifications:"
    )

    req_tbl = doc.add_table(rows=1, cols=3)
    req_widths = [1.8, 2.3, 2.4]
    req_headers = ["Resource Component", "Minimum Required", "Recommended Production"]
    req_data = [
        ["Operating System", "Windows 10 / Server 2019+, Ubuntu 20.04+, macOS 12+", "Windows 11 Pro / Enterprise 64-bit or Ubuntu 22.04 LTS"],
        ["Processor (CPU)", "4 x86_64 Cores (2.4 GHz+)", "8 Cores (Intel Core i7/i9 or AMD Ryzen 7/9 / Xeon)"],
        ["System Memory (RAM)", "8 GB DDR4", "16 GB DDR4/DDR5 (32 GB if running 8B LLM locally)"],
        ["Storage Space", "2.0 GB free SSD space", "10.0 GB free NVMe SSD (for logs, reports, models)"],
        ["Python Runtime", "Python 3.10, 3.11, 3.12, or 3.13", "Python 3.11 or 3.13 64-bit"],
        ["Network Adapter", "Loopback (127.0.0.1) only", "Gigabit Ethernet (for LAN control room access)"],
        ["Web Browser", "Chromium 90+, Edge 90+, Firefox 90+", "Latest Microsoft Edge or Google Chrome"]
    ]
    format_table(req_tbl, req_widths, req_headers, req_data, header_bg="0A2540")

    # ==========================================
    # SECTION 2: PYTHON ENVIRONMENT PREPARATION
    # ==========================================
    h1 = doc.add_heading("2. Python Environment & Virtual Environment Setup", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "To ensure complete dependency isolation from host system packages, Operation Advisor must be installed "
        "within a dedicated Python virtual environment."
    )

    h2 = doc.add_heading("2.1 Verifying Python Installation", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "Open your system terminal (PowerShell on Windows, or Bash on Linux/macOS) and verify that Python 64-bit is available in your PATH:"
    )
    add_code_block(doc, "python --version\n# Expected output: Python 3.10.x, 3.11.x, 3.12.x, or 3.13.x")

    add_callout(
        doc,
        "Windows PowerShell Script Execution Policy: If you encounter an execution policy error when activating "
        "the virtual environment in Windows PowerShell, run: \n"
        "Set-ExecutionPolicy -Scope Process -ExecutionPolicy Bypass",
        title="POWERSHELL PERMISSION NOTICE",
        callout_type="WARNING"
    )

    h2 = doc.add_heading("2.2 Creating and Activating the Virtual Environment", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph("Navigate to the Operation Advisor root project directory and initialize the virtual environment:")

    doc.add_paragraph("For Windows (PowerShell / Command Prompt):")
    add_code_block(doc, "cd C:\\Users\\salmb\\Desktop\\operation-advisor\npython -m venv venv_win\n.\\venv_win\\Scripts\\Activate.ps1")

    doc.add_paragraph("For Linux / macOS:")
    add_code_block(doc, "cd ~/operation-advisor\npython3 -m venv .venv\nsource .venv/bin/activate")

    # ==========================================
    # SECTION 3: DEPENDENCY INSTALLATION
    # ==========================================
    h1 = doc.add_heading("3. Package Dependency Installation", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Operation Advisor relies on a lightweight, curated set of industrial computing, probabilistic inference, and testing packages."
    )

    dep_tbl = doc.add_table(rows=1, cols=3)
    dep_widths = [1.8, 1.4, 3.3]
    dep_headers = ["Package Name", "Version", "System Purpose & Architecture Role"]
    dep_data = [
        ["numpy", ">= 1.24.0", "Core multidimensional array operations and numerical physics matrix math"],
        ["scipy", ">= 1.10.0", "Scientific optimization, integration, and statistical evaluation"],
        ["pandas", ">= 2.0.0", "Structured tabular telemetry dataframes and shift log serialization"],
        ["scikit-learn", ">= 1.2.0", "TF-IDF procedural document vectorization and offline classifiers"],
        ["pgmpy", ">= 0.1.20", "Exact Bayesian Belief Network structure and probabilistic inference"],
        ["networkx", ">= 3.0.0", "Diagnostic Knowledge Graph mapping symptoms to equipment safeguards"],
        ["python-docx", ">= 1.1.0", "Automated Word Document user and installation manual generation"],
        ["pytest", ">= 7.4.0", "Automated test suite framework and test harness runner"],
        ["pytest-cov", ">= 4.1.0", "Test execution coverage analysis and regulatory verification"]
    ]
    format_table(dep_tbl, dep_widths, dep_headers, dep_data, header_bg="1E293B")

    h2 = doc.add_heading("3.1 Installing via pip", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph("Run pip to install all required dependencies inside the active virtual environment:")
    add_code_block(doc, "pip install --upgrade pip\npip install numpy scipy pandas scikit-learn pgmpy networkx python-docx pytest pytest-cov")

    add_callout(
        doc,
        "Pre-built Wheel Support: All dependencies are distributed as pre-compiled binary wheels for Windows, Linux, and macOS. "
        "A full C/C++ compiler toolchain (such as Visual Studio C++ Build Tools or gcc) is NOT required for installation.",
        title="BINARY WHEEL AVAILABILITY",
        callout_type="TIP"
    )

    # ==========================================
    # SECTION 4: REPOSITORY STRUCTURE & VERIFICATION
    # ==========================================
    h1 = doc.add_heading("4. Directory Structure & Corpus Verification", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph("Verify that the repository contains the following core directory hierarchy:")

    dir_tbl = doc.add_table(rows=1, cols=2)
    dir_widths = [2.2, 4.3]
    dir_headers = ["Directory / File Path", "Contents & Functional Role"]
    dir_data = [
        ["src/simulation_bridge/", "Dynamic physics equations of state (process_model.py) and 6 fault injectors (scenarios.py)."],
        ["src/reasoning/", "Rule engines (forward_chaining.py, backward_chaining.py) and Plant Knowledge Graph (knowledge_graph.py)."],
        ["src/learning/", "Bayesian belief network (bayesian_network.py) and supervised scenario classifiers (classifier_trainer.py)."],
        ["src/planning/", "Bounded A* search planner (astar_planner.py) and genetic algorithm optimizer (ga_optimizer.py)."],
        ["src/cognitive/", "Natural language parser (nlu_parser.py), working memory (working_memory.py), and Ollama/RAG client (ollama_client.py)."],
        ["src/rag/", "Document passage indexer (document_indexer.py), citation generator (citation_generator.py), and PDF builder (pdf_generator.py)."],
        ["src/ui/", "HTTP server (server.py), application launcher (app.py), and responsive Cockpit GUI (index.html)."],
        ["data/documents/", "Indexed standard operating procedures (SOP-V101.md, SOP-K101.md) and hazard tables (HAZOP-SEP.md)."],
        ["tests/", "Comprehensive unit, integration, and end-to-end evaluation suite (74 automated tests)."]
    ]
    format_table(dir_tbl, dir_widths, dir_headers, dir_data, header_bg="0A2540")

    # ==========================================
    # SECTION 5: AUTOMATED ACCEPTANCE TESTING
    # ==========================================
    h1 = doc.add_heading("5. Automated Test Suite & Acceptance Testing", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Before launching the web service, execute the full automated test suite to verify mathematical accuracy, "
        "reasoning consistency, and anti-hallucination citation validation across all modules:"
    )
    add_code_block(doc, ".\\venv_win\\Scripts\\python.exe -m pytest tests/\n# On Linux/macOS: python3 -m pytest tests/")

    doc.add_paragraph(
        "Expected Verification Output:\n"
        "• Total Tests: 74 test cases executed across 10 test modules.\n"
        "• Passing Rate: 100% (74 passed, 0 failed).\n"
        "• Execution Latency: ~7 to 9 seconds on modern multi-core processors."
    )

    add_callout(
        doc,
        "Acceptance Criteria: Do NOT proceed to production deployment if any test fails. "
        "A passing test suite verifies that the Bayesian Net, A* search planner, Forward Chaining, "
        "and Citation Guard are operating within strict deterministic bounds.",
        title="REGULATORY ACCEPTANCE GATE",
        callout_type="IMPORTANT"
    )

    # ==========================================
    # SECTION 6: SERVICE LAUNCH & NETWORK CONFIGURATION
    # ==========================================
    h1 = doc.add_heading("6. Launching the Cockpit & Network Deployment", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    h2 = doc.add_heading("6.1 Starting the Web Server (Local Workstation)", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "To start the Operation Advisor Cockpit on the local workstation console:"
    )
    add_code_block(doc, ".\\venv_win\\Scripts\\python.exe src/ui/app.py --port 8080")

    doc.add_paragraph(
        "The server initializes on port 8080. Open any modern browser and navigate to:\n"
        "http://localhost:8080"
    )

    h2 = doc.add_heading("6.2 Network-Wide Control Room Binding (0.0.0.0)", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "To allow multiple operator consoles across the control room LAN to access the Cockpit, "
        "bind the host address to all interfaces:"
    )
    add_code_block(doc, ".\\venv_win\\Scripts\\python.exe src/ui/app.py --host 0.0.0.0 --port 8080")

    doc.add_paragraph(
        "Remote console operators can then access the interface via the workstation's static IP address (e.g. http://192.168.1.100:8080)."
    )

    h2 = doc.add_heading("6.3 Windows Firewall Configuration", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "If accessing from a remote terminal, allow inbound TCP traffic on port 8080 in Windows Defender Firewall via elevated PowerShell:"
    )
    add_code_block(doc, "New-NetFirewallRule -DisplayName \"Operation Advisor Cockpit\" -Direction Inbound -LocalPort 8080 -Protocol TCP -Action Allow")

    # ==========================================
    # SECTION 7: OPTIONAL LOCAL LLM (OLLAMA) INTEGRATION
    # ==========================================
    h1 = doc.add_heading("7. Optional Neural LLM Integration (Ollama)", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Operation Advisor includes a deterministic, zero-hallucination RAG synthesizer that operates completely offline without an LLM. "
        "However, if natural language generative conversational responses are desired, the system seamlessly interfaces with a local Ollama instance."
    )

    ollama_steps = [
        ("Step 1: Download & Install Ollama: ", "Obtain the Windows or Linux installer from https://ollama.com and complete installation."),
        ("Step 2: Pull the Recommended Model: ", "Open a terminal and download the lightweight, high-speed Llama 3.2 model:"),
        ("Step 3: Verify Ollama Service: ", "Ensure the service is listening on default port 11434 (http://127.0.0.1:11434)."),
        ("Step 4: Automatic Cockpit Detection: ", "Operation Advisor automatically detects the running Ollama instance upon startup, updating the top status badge to 'Ollama: llama3.2 (Active)'.")
    ]
    for o_title, o_desc in ollama_steps:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_b = bp.add_run(o_title)
        r_b.bold = True
        r_t = bp.add_run(o_desc)

    add_code_block(doc, "ollama pull llama3.2\n# Alternative models: ollama pull mistral or ollama pull llama3.1")

    # ==========================================
    # SECTION 8: SYSTEM TROUBLESHOOTING & MAINTENANCE
    # ==========================================
    h1 = doc.add_heading("8. Troubleshooting & Maintenance", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    faq_tbl = doc.add_table(rows=1, cols=3)
    faq_widths = [1.8, 2.2, 2.5]
    faq_headers = ["Symptom / Error", "Root Cause", "Corrective Action"]
    faq_data = [
        ["OSError: [Errno 10048] address already in use", "A prior server instance is still bound to port 8080.", "Stop the process using PowerShell:\nStop-Process -Id (Get-NetTCPConnection -LocalPort 8080).OwningProcess -Force"],
        ["ModuleNotFoundError: No module named 'docx'", "python-docx was not installed in the active virtual environment.", "Run: .\\venv_win\\Scripts\\pip.exe install python-docx"],
        ["Browser displays 'Failed to fetch' or blank screen", "The Cockpit server was stopped or firewall blocked port 8080.", "Restart server via app.py and verify port status using: Get-NetTCPConnection -LocalPort 8080"],
        ["LLM status shows 'Deterministic RAG'", "Ollama is either stopped or blocked by firewall.", "Start Ollama desktop service or continue using deterministic RAG (zero hallucination guaranteed)."]
    ]
    format_table(faq_tbl, faq_widths, faq_headers, faq_data, header_bg="1E293B")

    # Final signoff
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- End of Operation Advisor Installation & Deployment Guide (Document Ref: INS-OA-TR1-2026) ---")
    r_end.font.size = Pt(9)
    r_end.font.color.rgb = RGBColor(148, 163, 184)

    # Save to docs and root
    os.makedirs("docs", exist_ok=True)
    doc_path = os.path.join("docs", "Operation_Advisor_Installation_Guide.docx")
    root_path = "INSTALLATION_GUIDE.docx"
    doc.save(doc_path)
    doc.save(root_path)
    print(f"Installation guide successfully generated at:\n  - {doc_path}\n  - {root_path}")

if __name__ == "__main__":
    build_installation_guide()

